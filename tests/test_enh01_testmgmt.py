"""
Tests for testmgmt/ -- the new, additive Post-MVP2 Enhancement 01 Excel
test-management workbook generator. Verifies the workbooks are a real,
deterministic, faithful view over the real, persisted MVP2 corpus -- never
a second, hand-editable source of truth.
"""
from __future__ import annotations

from openpyxl import load_workbook

from testmgmt.generate_workbooks import build_all_workbooks


def test_build_all_workbooks_produces_three_real_files():
    paths = build_all_workbooks()
    assert set(paths.keys()) == {"testcases", "testdata", "traceability"}
    for path in paths.values():
        assert path.exists()
        assert path.stat().st_size > 0


def test_testcase_workbook_contains_every_real_persisted_testcase():
    paths = build_all_workbooks()
    wb = load_workbook(paths["testcases"])
    ws = wb["Testcases"]
    ids_in_sheet = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}

    import testcases.persist as testcases_persist

    real_ids = set(testcases_persist.list_persisted_testcase_ids())
    assert ids_in_sheet == real_ids


def test_testcase_workbook_never_rewrites_the_disclosed_generic_steps():
    """The historical testcase-content limitation must be shown exactly as
    persisted, never rephrased to look more complete."""
    paths = build_all_workbooks()
    wb = load_workbook(paths["testcases"])
    ws = wb["Testcases"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "TC-REQ-REG-01-01":
            steps_cell = ws.cell(row=row, column=6).value
            assert "Perform the action described by the approved requirement." in steps_cell
            assert "Observe the resulting system behavior." in steps_cell


def test_testdata_workbook_has_a_row_per_real_field():
    paths = build_all_workbooks()
    wb = load_workbook(paths["testdata"])
    ws = wb["TestData Fields"]

    import testdata.persist as testdata_persist

    total_fields = 0
    for did in testdata_persist.list_persisted_dataset_ids():
        ds = testdata_persist.load_persisted_dataset(did)
        total_fields += len(ds.get("fields", []))

    assert ws.max_row - 1 == total_fields


def test_testdata_workbook_mapping_sheet_supports_one_testcase_many_datasets():
    paths = build_all_workbooks()
    wb = load_workbook(paths["testdata"])
    mapping = wb["Testcase-Dataset Mapping"]
    tc_counts = {}
    for row in range(2, mapping.max_row + 1):
        tc_id = mapping.cell(row=row, column=1).value
        tc_counts[tc_id] = tc_counts.get(tc_id, 0) + 1
    assert tc_counts.get("TC-REQ-REG-01-01", 0) == 6


def test_traceability_workbook_reflects_real_persisted_coverage():
    paths = build_all_workbooks()
    wb = load_workbook(paths["traceability"])
    ws = wb["Req-Testcase Traceability"]
    req_ids = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert req_ids == {"REQ-ACO-03", "REQ-CART-03", "REQ-PAY-01", "REQ-REG-01"}


def test_workbook_regeneration_is_deterministic_in_content():
    """Regenerating against an unchanged corpus must not silently drift."""
    paths_a = build_all_workbooks()
    wb_a = load_workbook(paths_a["testcases"])
    values_a = [tuple(r) for r in wb_a["Testcases"].iter_rows(values_only=True)]

    paths_b = build_all_workbooks()
    wb_b = load_workbook(paths_b["testcases"])
    values_b = [tuple(r) for r in wb_b["Testcases"].iter_rows(values_only=True)]

    assert values_a == values_b
